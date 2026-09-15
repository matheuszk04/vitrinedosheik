"""Traz a planilha de enriquecimento preenchida de volta para o catálogo.

    python3 src/import-enrichment.py enriquecimento.xlsx [--aplicar]

Só grava o que estiver preenchido. Célula vazia não apaga dado existente.
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CATALOG = ROOT / "content" / "perfumes.json"

SPLIT_FIELDS = {
    "familias": "families",
    "personalidade": "personality",
    "ocasioes": "occasions",
    "momentos": "moments",
    "climas": "climates",
}
NOTE_FIELDS = {"notas_topo": "top", "notas_coracao": "heart", "notas_fundo": "base"}
PROFILE_FIELDS = {"intensidade": "intensity", "projecao": "projection", "fixacao": "longevity"}
PLAIN_FIELDS = {"marca": "brand", "genero": "gender", "volume_ml": "volumeMl", "concentracao": "concentration"}


def split(value):
    if value is None:
        return None
    parts = [p.strip() for p in str(value).split(";") if p.strip()]
    return parts or None


def clean(value):
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def main():
    if len(sys.argv) < 2:
        sys.exit("uso: python3 src/import-enrichment.py enriquecimento.xlsx [--aplicar]")
    apply = "--aplicar" in sys.argv

    ws = openpyxl.load_workbook(sys.argv[1], data_only=True).active
    header = [str(c.value).strip() if c.value else "" for c in ws[2]]
    idx = {name: i for i, name in enumerate(header)}

    perfumes = json.loads(CATALOG.read_text(encoding="utf8"))
    by_slug = {p["slug"]: p for p in perfumes}

    touched, unknown, now_ready = [], [], []

    for row in ws.iter_rows(min_row=3, values_only=True):
        slug = clean(row[idx["slug"]]) if "slug" in idx else None
        if not slug:
            continue
        p = by_slug.get(slug)
        if not p:
            unknown.append(slug)
            continue

        was_ready = bool(p.get("families") and p.get("personality") and p.get("profile"))
        changes = []

        for col, field in PLAIN_FIELDS.items():
            v = clean(row[idx[col]]) if col in idx else None
            if v and str(p.get(field) or "") != v:
                p[field] = int(v) if field == "volumeMl" and v.isdigit() else v
                changes.append(field)

        for col, field in SPLIT_FIELDS.items():
            v = split(row[idx[col]]) if col in idx else None
            if v and p.get(field) != v:
                p[field] = v
                changes.append(field)

        notes = dict(p.get("notes") or {})
        for col, key in NOTE_FIELDS.items():
            v = split(row[idx[col]]) if col in idx else None
            if v and notes.get(key) != v:
                notes[key] = v
                changes.append(f"notes.{key}")
        if notes:
            p["notes"] = notes

        profile = dict(p.get("profile") or {})
        for col, key in PROFILE_FIELDS.items():
            v = clean(row[idx[col]]) if col in idx else None
            if v and str(v).isdigit():
                n = max(1, min(5, int(v)))
                if profile.get(key) != n:
                    profile[key] = n
                    changes.append(f"profile.{key}")
        if profile:
            p["profile"] = profile

        if changes:
            touched.append((p["name"], changes))
            is_ready = bool(p.get("families") and p.get("personality") and p.get("profile"))
            if is_ready and not was_ready:
                now_ready.append(p["name"])

    print(f"{len(touched)} perfumes com alteração\n")
    for name, changes in touched:
        print(f"   ~ {name}: {', '.join(changes)}")
    if now_ready:
        print(f"\n✓ passaram a ter perfil completo ({len(now_ready)}):")
        for n in now_ready:
            print(f"   + {n}")
    if unknown:
        print(f"\n⚠ slug não encontrado no catálogo: {', '.join(unknown)}")

    if apply:
        CATALOG.write_text(json.dumps(perfumes, ensure_ascii=False, indent=2) + "\n", encoding="utf8")
        print(f"\n✓ {CATALOG.relative_to(ROOT)} atualizado")
    else:
        print("\n(simulação — rode com --aplicar para gravar)")


if __name__ == "__main__":
    main()
