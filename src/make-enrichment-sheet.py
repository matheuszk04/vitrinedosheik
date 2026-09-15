"""Gera a planilha de enriquecimento a partir do catálogo atual.

    python3 src/make-enrichment-sheet.py

Produz enriquecimento.xlsx: uma linha por perfume, com o que já existe
preenchido. O lojista completa o resto e devolve via import-enrichment.py.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
CATALOG = ROOT / "content" / "perfumes.json"
OUT = ROOT / "enriquecimento.xlsx"

# (coluna, largura, essencial para a descoberta funcionar?, ajuda)
COLUMNS = [
    ("slug", 20, False, "NÃO EDITE — liga a linha ao perfume no site"),
    ("nome", 26, False, "NÃO EDITE"),
    ("marca", 16, True, "Lattafa, Armaf, Afnan, Rasasi..."),
    ("genero", 14, True, "masculino, feminino ou unissex"),
    ("familias", 30, True, "separe por ; — Oriental;Amadeirado;Especiado"),
    ("personalidade", 30, True, "3 palavras, separe por ; — Noturno;Intenso;Magnético"),
    ("ocasioes", 30, True, "separe por ; — Noite;Encontro;Evento;Trabalho;Dia"),
    ("momentos", 22, True, "separe por ; — manhã;tarde;noite;madrugada"),
    ("climas", 20, True, "separe por ; — quente;ameno;frio"),
    ("intensidade", 13, True, "1 a 5"),
    ("projecao", 12, True, "1 a 5 — o quanto os outros sentem"),
    ("fixacao", 11, True, "1 a 5 — quanto tempo dura"),
    ("notas_topo", 30, False, "separe por ; — o que se sente nos primeiros minutos"),
    ("notas_coracao", 30, False, "separe por ;"),
    ("notas_fundo", 30, False, "separe por ;"),
    ("volume_ml", 11, False, "só o número: 100"),
    ("concentracao", 14, False, "EDP, EDT, Extrait..."),
]

GOLD = "C9A45C"
DARK = "1A1A1A"


def joined(value):
    return ";".join(value) if isinstance(value, list) else ""


def main():
    perfumes = json.loads(CATALOG.read_text(encoding="utf8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Enriquecimento"

    # linha 1: ajuda · linha 2: cabeçalho
    for i, (name, width, essential, hint) in enumerate(COLUMNS, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width

        help_cell = ws.cell(row=1, column=i, value=hint)
        help_cell.font = Font(size=8, italic=True, color="888888")
        help_cell.alignment = Alignment(wrap_text=True, vertical="top")

        head = ws.cell(row=2, column=i, value=name)
        head.font = Font(bold=True, size=10, color=GOLD if essential else "DDDDDD")
        head.fill = PatternFill("solid", fgColor=DARK)
        head.alignment = Alignment(horizontal="center")

    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "C3"

    for r, p in enumerate(perfumes, 3):
        values = {
            "slug": p.get("slug"),
            "nome": p.get("name"),
            "marca": p.get("brand"),
            "genero": p.get("gender"),
            "familias": joined(p.get("families")),
            "personalidade": joined(p.get("personality")),
            "ocasioes": joined(p.get("occasions")),
            "momentos": joined(p.get("moments")),
            "climas": joined(p.get("climates")),
            "intensidade": (p.get("profile") or {}).get("intensity"),
            "projecao": (p.get("profile") or {}).get("projection"),
            "fixacao": (p.get("profile") or {}).get("longevity"),
            "notas_topo": joined((p.get("notes") or {}).get("top")),
            "notas_coracao": joined((p.get("notes") or {}).get("heart")),
            "notas_fundo": joined((p.get("notes") or {}).get("base")),
            "volume_ml": p.get("volumeMl"),
            "concentracao": p.get("concentration"),
        }
        for c, (name, *_rest) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=r, column=c, value=values.get(name))
            if name in ("slug", "nome"):
                cell.font = Font(color="777777")

    last = len(perfumes) + 2
    gender_dv = DataValidation(type="list", formula1='"masculino,feminino,unissex"', allow_blank=True)
    ws.add_data_validation(gender_dv)
    gender_dv.add(f"D3:D{last}")

    scale_dv = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    ws.add_data_validation(scale_dv)
    for col in ("J", "K", "L"):
        scale_dv.add(f"{col}3:{col}{last}")

    wb.save(OUT)

    essentials = [c[0] for c in COLUMNS if c[2]]
    faltando = sum(
        1 for p in perfumes
        if not (p.get("families") and p.get("personality") and p.get("profile"))
    )
    print(f"✓ {OUT.name}: {len(perfumes)} perfumes")
    print(f"  {faltando} ainda sem perfil para recomendação")
    print(f"  colunas em dourado são as que a descoberta precisa: {', '.join(essentials)}")


if __name__ == "__main__":
    main()
